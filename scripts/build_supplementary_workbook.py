"""Build formatted machine-readable supplementary tables workbook."""

from __future__ import annotations

import argparse
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


def format_sheet(sheet):
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = fill
    for column in range(1, sheet.max_column + 1):
        values = [str(sheet.cell(row, column).value or "") for row in range(1, min(sheet.max_row, 200) + 1)]
        width = min(45, max(10, max(len(value) for value in values) + 2))
        sheet.column_dimensions[get_column_letter(column)].width = width


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", type=Path, required=True)
    args = parser.parse_args()
    root = args.root.resolve()
    tables = root / "tables_q1"
    output = tables / "supplementary_tables.xlsx"
    sources = [
        ("S1_hypothesis_tests", "supp_table1_hypothesis_tests.csv"),
        ("S2_reciprocal_transfer", "supp_table2_reciprocal_transfer.csv"),
        ("S3_full_domain", "table5_full_domain_models.csv"),
        ("S4_scaffold", "supp_table4_scaffold.csv"),
        ("S5_crossed_bootstrap", "supp_table5_crossed_bootstrap.csv"),
        ("S6_ood_safeguards", "supp_table6_ood_safeguards.csv"),
        ("S6_ood_associations", "supp_table7_ood_associations.csv"),
        ("S7_source_winners", "supp_table7_fair_tuning_winners.csv"),
        ("S7_external", "supp_table7_extended_external.csv"),
        ("S8_dataset_flow", "table1_dataset_flow.csv"),
        ("S9_historical_selection", "table3_source_selection.csv"),
    ]
    repairs = [
        ("S8_historical_gaps", root / "results/repair_historical_population/gaps.csv"),
        ("S9_stage_populations", root / "results/repair_historical_population/stage_populations.csv"),
        ("S10_Holm_correction", root / "results/repair_holm/corrected_comparison.csv"),
        ("S4_matched_contrast", root / "results/repair_scaffold_comparison/summary.csv"),
        ("S11_descriptor_guard", root / "results/repair_descriptor_sensitivity/ensemble_metrics.csv"),
    ]
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        for sheet, filename in sources:
            frame = pd.read_csv(tables / filename)
            frame.to_excel(writer, sheet_name=sheet, index=False)
            format_sheet(writer.book[sheet])
        for sheet, path in repairs:
            pd.read_csv(path).to_excel(writer, sheet_name=sheet, index=False)
            format_sheet(writer.book[sheet])
        trace = pd.read_csv(root / "results/repair_descriptor_sensitivity/source_trace.csv")
        trace = trace.loc[trace.drug_id == "SMI_48082"]
        trace.to_excel(writer, sheet_name="S11_feature_trace", index=False)
        format_sheet(writer.book["S11_feature_trace"])
        metadata = pd.DataFrame(
            {
                "Field": ["Project", "Primary endpoint", "Bootstrap", "Bootstrap seed", "Generated from"],
                "Value": [
                    "Cross-screen cancer drug-response benchmark",
                    "Equal-drug-weighted macro Spearman",
                    "BCa, 9,999 drug-cluster resamples",
                    "20261004",
                    "Final saved result CSV files",
                ],
            }
        )
        metadata.to_excel(writer, sheet_name="README", index=False)
        format_sheet(writer.book["README"])
    print(output, output.stat().st_size)


if __name__ == "__main__":
    main()
